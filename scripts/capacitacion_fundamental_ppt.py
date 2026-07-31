"""Generador de la capacitación de Análisis Fundamental en PowerPoint.

Producido por el Área de Estudios y Post-Venta para el equipo comercial e IBS.

Enfoque editorial (definido por el director): el eje es **conceptual y educativo** —
que quien la reciba aprenda a *interpretar* cada publicación de datos, no a memorizar
reacciones de precio. La implicación de precio se incluye siempre, pero como
consecuencia razonada del concepto, nunca como recetario.

Tipografía: se usan Segoe UI y Consolas en lugar de las fuentes de marca
(Syne / DM Sans / Space Grotesk) porque estas solo existen en el repo como .woff2
—formato web— y no están instaladas en los equipos. PowerPoint las sustituiría y
rompería la maqueta en el PC de cada destinatario. Consolas mantiene el criterio de
cifras monoespaciadas del kit (equivalente a `tabular-nums`).

Uso:
    uv run --with python-pptx python scripts/capacitacion_fundamental_ppt.py
    uv run --with python-pptx python scripts/capacitacion_fundamental_ppt.py --out ruta.pptx
"""

from __future__ import annotations

import argparse
import math
import re
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------------------
# Kit de marca
# --------------------------------------------------------------------------------------

FONDO = RGBColor(0x0D, 0x0D, 0x1A)
SUPERFICIE = RGBColor(0x16, 0x16, 0x2A)
SUPERFICIE_ALT = RGBColor(0x1E, 0x1E, 0x36)
VERDE = RGBColor(0x00, 0xDC, 0x82)
ROJO = RGBColor(0xE8, 0x40, 0x40)
AMBAR = RGBColor(0xF5, 0xA5, 0x24)
AZUL = RGBColor(0x4C, 0x8D, 0xFF)
GRIS = RGBColor(0xA9, 0xA5, 0xB4)
GRIS_TENUE = RGBColor(0x6E, 0x6A, 0x7C)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

TIPO = "Segoe UI"
TIPO_CIFRA = "Consolas"

ANCHO = Inches(13.333)
ALTO = Inches(7.5)
MARGEN = Inches(0.72)
UTIL = ANCHO - 2 * MARGEN

# Retrato del autor, ya recortado en círculo. Lo fija `construir()` a partir del
# argumento `--foto`; si no se entrega, portada y cierre caen al diseño sin retrato.
FOTO: Path | None = None

# --------------------------------------------------------------------------------------
# Medición de texto
#
# PowerPoint no expone métricas de fuente, así que el alto de cada bloque se estima
# antes de dibujar. Sin esto el layout se rompe de dos maneras que ya vimos: un título
# que ocupa dos líneas se superpone con el párrafo siguiente, y una tabla larga se
# expande por debajo del pie de página. Los factores están calibrados contra renders
# reales de Segoe UI a 1600×900.
# --------------------------------------------------------------------------------------

CHARS_POR_PULGADA = 148.0   # a 1 pt de tamaño; texto normal
CHARS_POR_PULGADA_NEG = 138.0  # la negrita ocupa más ancho por carácter


def _pulg(emu) -> float:
    return float(emu) / 914400.0


def _limpio(texto: str) -> str:
    return texto.replace("*", "")


def _n_lineas(texto: str, ancho_in: float, tam: float, negrita: bool = False) -> int:
    """Cuántas líneas ocupará `texto` en un ancho dado. Respeta los saltos explícitos."""
    if not texto:
        return 1
    factor = CHARS_POR_PULGADA_NEG if negrita else CHARS_POR_PULGADA
    por_linea = max(8.0, ancho_in * factor / tam)
    return sum(
        max(1, math.ceil(len(_limpio(segmento)) / por_linea))
        for segmento in texto.split("\n")
    )


def _alto_texto(lineas, ancho_in: float, tam: float, interlineado: float = 1.3,
                espacio_antes: float = 6.0, negrita: bool = False) -> float:
    """Alto en pulgadas que necesita un bloque de párrafos."""
    items = [lineas] if isinstance(lineas, str) else list(lineas)
    alto_linea = tam * interlineado / 72.0
    total = 0.0
    for indice, linea in enumerate(items):
        total += _n_lineas(linea, ancho_in, tam, negrita) * alto_linea
        if indice:
            total += espacio_antes / 72.0
    return total


# Alto fijo de una caja titulada: filete + título + separación + margen inferior.
CAJA_CROMO = 0.78


def _alto_caja(lineas, ancho_caja_in: float, tam: float,
               interlineado: float = 1.3) -> float:
    return CAJA_CROMO + _alto_texto(lineas, ancho_caja_in - 0.4, tam, interlineado)


def _tam_que_cabe(medir, disponible: float, tam_inicial: float,
                  minimo: float = 8.5, paso: float = 0.5) -> float:
    """Reduce el tamaño de fuente hasta que el bloque quepa en `disponible`."""
    tam = tam_inicial
    while tam > minimo and medir(tam) > disponible:
        tam -= paso
    return tam


# --------------------------------------------------------------------------------------
# Primitivas de dibujo
# --------------------------------------------------------------------------------------


def _fondo(slide, color=FONDO) -> None:
    relleno = slide.background.fill
    relleno.solid()
    relleno.fore_color.rgb = color


def _rect(slide, left, top, width, height, relleno=None, borde=None, grosor=1.0,
          forma=MSO_SHAPE.RECTANGLE):
    figura = slide.shapes.add_shape(forma, left, top, width, height)
    if relleno is None:
        figura.fill.background()
    else:
        figura.fill.solid()
        figura.fill.fore_color.rgb = relleno
    if borde is None:
        figura.line.fill.background()
    else:
        figura.line.color.rgb = borde
        figura.line.width = Pt(grosor)
    figura.shadow.inherit = False
    figura.text_frame.text = ""
    return figura


def _runs(parrafo, texto: str, tam: float, color, tipo: str = TIPO,
          color_fuerte=None, negrita_base: bool = False) -> None:
    """Escribe `texto` en el párrafo interpretando *asterisco* como destacado."""
    for indice, parte in enumerate(re.split(r"\*(.+?)\*", texto)):
        if not parte:
            continue
        destacado = indice % 2 == 1
        run = parrafo.add_run()
        run.text = parte
        run.font.size = Pt(tam)
        run.font.name = tipo
        run.font.bold = destacado or negrita_base
        run.font.color.rgb = (color_fuerte or BLANCO) if destacado else color


def _texto(slide, left, top, width, height, lineas, tam=14, color=GRIS, tipo=TIPO,
           alineacion=PP_ALIGN.LEFT, anclaje=MSO_ANCHOR.TOP, interlineado=1.24,
           espacio_antes=0, color_fuerte=None, negrita=False):
    """`lineas` acepta un str o una lista de str (un párrafo por elemento)."""
    caja = slide.shapes.add_textbox(left, top, width, height)
    marco = caja.text_frame
    marco.word_wrap = True
    marco.vertical_anchor = anclaje
    marco.margin_left = marco.margin_right = 0
    marco.margin_top = marco.margin_bottom = 0

    items = [lineas] if isinstance(lineas, str) else list(lineas)
    for indice, linea in enumerate(items):
        parrafo = marco.paragraphs[0] if indice == 0 else marco.add_paragraph()
        parrafo.alignment = alineacion
        parrafo.line_spacing = interlineado
        if indice > 0 and espacio_antes:
            parrafo.space_before = Pt(espacio_antes)
        _runs(parrafo, linea, tam, color, tipo, color_fuerte, negrita)
    return caja


def _encabezado(slide, titulo: str, bajada: str | None = None, acento=VERDE,
                etiqueta: str | None = None):
    """Título de slide con filete de acento. Devuelve el borde superior del cuerpo."""
    tope = Inches(0.52)
    if etiqueta:
        _texto(slide, MARGEN, tope, UTIL, Inches(0.26), etiqueta.upper(),
               tam=10.5, color=acento, negrita=True)
        tope = tope + Inches(0.3)

    alto_titulo = _alto_texto(titulo, _pulg(UTIL), 29, 1.02, negrita=True)
    _texto(slide, MARGEN, tope, UTIL, Inches(alto_titulo), titulo,
           tam=29, color=BLANCO, negrita=True, interlineado=1.02)
    borde = tope + Inches(alto_titulo + 0.14)

    _rect(slide, MARGEN, borde, Inches(1.05), Emu(31750), relleno=acento)
    borde = borde + Inches(0.2)

    if bajada:
        alto_bajada = _alto_texto(bajada, _pulg(UTIL), 13.5, 1.26)
        _texto(slide, MARGEN, borde, UTIL, Inches(alto_bajada), bajada,
               tam=13.5, color=GRIS, interlineado=1.26, color_fuerte=acento)
        borde = borde + Inches(alto_bajada)
    return borde + Inches(0.26)


def _pie(slide, numero: int, seccion: str = "") -> None:
    y = ALTO - Inches(0.5)
    izquierda = "Área de Estudios y Post-Venta · GI"
    if seccion:
        izquierda += f"  ·  {seccion}"
    _texto(slide, MARGEN, y, UTIL - Inches(0.8), Inches(0.24), izquierda,
           tam=8.5, color=GRIS_TENUE)
    _texto(slide, ANCHO - MARGEN - Inches(0.8), y, Inches(0.8), Inches(0.24),
           str(numero), tam=8.5, color=GRIS_TENUE, tipo=TIPO_CIFRA,
           alineacion=PP_ALIGN.RIGHT)


EXTENSIONES_FOTO = (".png", ".jpg", ".jpeg", ".webp")


def _resolver_foto(ruta: Path | None) -> Path | None:
    """Ubica el retrato. Si la ruta exacta no existe, toma cualquier imagen de la
    carpeta: el archivo suele llegar con el nombre que le pone WhatsApp."""
    if ruta is None:
        return None
    if ruta.exists():
        return ruta
    carpeta = ruta.parent
    if not carpeta.is_dir():
        return None
    candidatas = sorted(
        archivo for archivo in carpeta.iterdir()
        if archivo.suffix.lower() in EXTENSIONES_FOTO
        and not archivo.name.startswith(".")
    )
    if candidatas:
        print(f"  retrato: se usa {candidatas[0].name}")
        return candidatas[0]
    return None


def _foto_circular(origen: Path, lado: int = 720, zoom: float = 1.0,
                   centro_y: float = 0.5) -> Path | None:
    """Recorta la foto a un círculo y la deja en un PNG temporal.

    PowerPoint no aplica máscaras: para que el retrato salga redondo hay que
    recortarlo antes. `zoom` cierra el encuadre sobre el rostro —un círculo
    inscrito en la foto completa deja demasiado fondo compitiendo con la cara— y
    `centro_y` fija en qué fracción de la altura queda el centro del recorte.
    """
    try:
        from PIL import Image, ImageDraw  # noqa: PLC0415
    except ImportError:
        print("  aviso: Pillow no disponible, la foto se omite")
        return None

    imagen = Image.open(origen).convert("RGBA")
    corto = int(min(imagen.size) / max(1.0, zoom))
    izq = max(0, (imagen.width - corto) // 2)
    arriba = int(imagen.height * centro_y - corto / 2)
    arriba = max(0, min(arriba, imagen.height - corto))
    imagen = imagen.crop((izq, arriba, izq + corto, arriba + corto))
    imagen = imagen.resize((lado, lado), Image.LANCZOS)

    mascara = Image.new("L", (lado * 4, lado * 4), 0)
    ImageDraw.Draw(mascara).ellipse((0, 0, lado * 4 - 1, lado * 4 - 1), fill=255)
    imagen.putalpha(mascara.resize((lado, lado), Image.LANCZOS))

    aro = ImageDraw.Draw(imagen)
    grosor = max(2, lado // 120)
    aro.ellipse((grosor // 2, grosor // 2, lado - grosor // 2, lado - grosor // 2),
                outline=(0, 220, 130, 255), width=grosor)

    # Al temporal del sistema y no junto al original: es un intermedio derivado y no
    # tiene por qué ensuciar (ni acabar versionado en) la carpeta de assets.
    destino = Path(tempfile.gettempdir()) / f"gi_retrato_{origen.stem}.png"
    imagen.save(destino, "PNG")
    return destino


def _bloque_autor(slide, left, top, autor: dict, foto: Path | None,
                  diametro: float = 1.32, tam_nombre: float = 14.0,
                  compacto: bool = False):
    """Retrato circular + nombre y credenciales. Devuelve el alto ocupado.

    `compacto` usa la variante de dos líneas de las credenciales, para el cierre,
    donde el espacio restante bajo los dos bloques de advertencia es escaso.
    """
    lineas_cred = len(autor["credenciales_compactas" if compacto else "credenciales"])
    alto_texto = 0.6 + lineas_cred * 0.21

    x_texto = left
    desplazo = 0.0
    if foto is not None:
        slide.shapes.add_picture(str(foto), left, top, Inches(diametro),
                                 Inches(diametro))
        x_texto = left + Inches(diametro + 0.28)
        # El retrato es más alto que el texto: se centra el texto contra el círculo
        # en vez de dejarlo colgando del borde superior.
        desplazo = max(0.0, (diametro - alto_texto) / 2)

    cursor = top + Inches(0.06 + desplazo)
    _texto(slide, x_texto, cursor, Inches(7.4), Inches(0.2),
           autor.get("rol", "Preparado por").upper(), tam=8.5, color=GRIS_TENUE,
           negrita=True)
    cursor = cursor + Inches(0.24)
    _texto(slide, x_texto, cursor, Inches(7.4), Inches(0.3), autor["nombre"],
           tam=tam_nombre, color=BLANCO, negrita=True)
    cursor = cursor + Inches(0.3)
    clave = "credenciales_compactas" if compacto else "credenciales"
    for linea in autor[clave]:
        _texto(slide, x_texto, cursor, Inches(7.4), Inches(0.22), linea,
               tam=10, color=GRIS, interlineado=1.2)
        cursor = cursor + Inches(0.21)
    return max(_pulg(cursor - top), diametro if foto is not None else 0.0)


def _chip(slide, left, top, etiqueta: str, valor: str, ancho, color=AZUL):
    """Par etiqueta/valor sobre superficie, para metadata de un indicador."""
    alto = Inches(0.62)
    _rect(slide, left, top, ancho, alto, relleno=SUPERFICIE)
    _rect(slide, left, top, Emu(28575), alto, relleno=color)
    _texto(slide, left + Inches(0.16), top + Inches(0.09), ancho - Inches(0.26),
           Inches(0.2), etiqueta.upper(), tam=8, color=GRIS_TENUE, negrita=True)
    _texto(slide, left + Inches(0.16), top + Inches(0.29), ancho - Inches(0.26),
           Inches(0.26), valor, tam=11, color=BLANCO, negrita=True, interlineado=1.05)


def _caja_titulada(slide, left, top, width, height, titulo: str, lineas,
                   color=VERDE, tam=12, relleno=SUPERFICIE):
    _rect(slide, left, top, width, height, relleno=relleno)
    _rect(slide, left, top, width, Emu(28575), relleno=color)
    _texto(slide, left + Inches(0.2), top + Inches(0.16), width - Inches(0.4),
           Inches(0.26), titulo.upper(), tam=9.5, color=color, negrita=True)
    _texto(slide, left + Inches(0.2), top + Inches(0.48), width - Inches(0.4),
           height - Inches(0.62), lineas, tam=tam, color=GRIS, interlineado=1.3,
           espacio_antes=5)


# --------------------------------------------------------------------------------------
# Tipos de slide
# --------------------------------------------------------------------------------------


def slide_portada(prs, datos, numero):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    _rect(slide, Emu(0), Emu(0), Inches(0.14), ALTO, relleno=VERDE)

    _texto(slide, MARGEN, Inches(1.35), UTIL, Inches(0.3), datos["kicker"].upper(),
           tam=12, color=VERDE, negrita=True)
    _texto(slide, MARGEN, Inches(1.85), Inches(10.6), Inches(1.9), datos["titulo"],
           tam=54, color=BLANCO, negrita=True, interlineado=0.98)
    _rect(slide, MARGEN, Inches(3.95), Inches(1.6), Emu(38100), relleno=VERDE)
    _texto(slide, MARGEN, Inches(4.25), Inches(9.6), Inches(1.0), datos["bajada"],
           tam=16, color=GRIS, interlineado=1.32, color_fuerte=BLANCO)

    # Autoría: el retrato y las credenciales van en la portada, no solo al cierre.
    # Es material que circula entre equipos, así que quién lo firma se ve de entrada.
    y = Inches(5.32)
    _rect(slide, MARGEN, y - Inches(0.3), UTIL, Emu(12700), relleno=SUPERFICIE_ALT)
    # 1,55" y no 1,32": a 1,32 el rostro quedaba demasiado pequeño para reconocerse
    # en la slide, y cerrar más el recorte cortaba la cabeza contra el círculo.
    DIAMETRO_PORTADA = 1.55
    _bloque_autor(slide, MARGEN, y, datos["autor"], FOTO,
                  diametro=DIAMETRO_PORTADA)

    left = MARGEN + Inches(9.35)
    for indice, (etiqueta, valor) in enumerate(datos["meta"]):
        tope = y + Inches(0.22) + indice * Inches(0.62)
        _texto(slide, left, tope, Inches(2.6), Inches(0.2), etiqueta.upper(),
               tam=8.5, color=GRIS_TENUE, negrita=True)
        _texto(slide, left, tope + Inches(0.22), Inches(2.6), Inches(0.26), valor,
               tam=11, color=BLANCO, negrita=True)
    return slide


def slide_seccion(prs, datos, numero):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    color = datos.get("color", VERDE)
    _rect(slide, Emu(0), Emu(0), Inches(0.14), ALTO, relleno=color)

    # El ancho se queda corto respecto del inicio del título (1,9") a propósito: así
    # las cajas no se cruzan y el verificador no reporta un solapamiento inexistente.
    _texto(slide, MARGEN, Inches(2.15), Inches(1.75), Inches(1.5), datos["numero"],
           tam=96, color=color, tipo=TIPO_CIFRA, negrita=True, interlineado=0.9)
    _texto(slide, MARGEN + Inches(1.9), Inches(2.35), Inches(9.0), Inches(1.1),
           datos["titulo"], tam=40, color=BLANCO, negrita=True, interlineado=1.0)
    _texto(slide, MARGEN + Inches(1.95), Inches(3.55), Inches(8.6), Inches(1.2),
           datos["bajada"], tam=14.5, color=GRIS, interlineado=1.34,
           color_fuerte=color)
    _pie(slide, numero)
    return slide


def slide_bullets(prs, datos, numero):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    acento = datos.get("color", VERDE)
    y = _encabezado(slide, datos["titulo"], datos.get("bajada"), acento,
                    datos.get("etiqueta"))

    clave = datos.get("clave")
    alto_clave = (_alto_caja(clave[1], _pulg(UTIL), 12.5) + 0.1) if clave else 0.0
    disponible = _pulg(ALTO - y) - 0.62 - alto_clave

    columnas = datos.get("columnas", 1)
    puntos = datos["puntos"]
    if columnas == 2:
        mitad = (len(puntos) + 1) // 2
        grupos = [puntos[:mitad], puntos[mitad:]]
        ancho_col = (UTIL - Inches(0.5)) / 2
    else:
        grupos = [puntos]
        ancho_col = UTIL

    normalizados = [
        [p if isinstance(p, tuple) else (None, p) for p in grupo] for grupo in grupos
    ]
    ancho_texto = _pulg(ancho_col) - 0.18

    def _alto_columna(tam: float, grupo) -> float:
        total = 0.0
        for cabeza, cuerpo in grupo:
            if cabeza:
                total += _alto_texto(cabeza, ancho_texto, tam + 0.5, 1.1,
                                     negrita=True) + 0.06
            total += _alto_texto(cuerpo, ancho_texto, tam, 1.28) + 0.16
        return total

    tam = _tam_que_cabe(
        lambda t: max(_alto_columna(t, g) for g in normalizados), disponible, 12.5, 9.5)

    for indice_col, grupo in enumerate(normalizados):
        left = MARGEN + indice_col * (ancho_col + Inches(0.5))
        cursor = y
        for cabeza, cuerpo in grupo:
            if cabeza:
                alto = _alto_texto(cabeza, ancho_texto, tam + 0.5, 1.1, negrita=True)
                _rect(slide, left, cursor + Inches(0.07), Emu(22225), Inches(0.19),
                      relleno=acento)
                _texto(slide, left + Inches(0.18), cursor, ancho_col - Inches(0.18),
                       Inches(alto), cabeza, tam=tam + 0.5, color=BLANCO,
                       negrita=True, interlineado=1.1)
                cursor = cursor + Inches(alto + 0.06)
            alto = _alto_texto(cuerpo, ancho_texto, tam, 1.28)
            desplazo = Inches(0.18) if cabeza else Emu(0)
            _texto(slide, left + desplazo, cursor, ancho_col - desplazo,
                   Inches(alto), cuerpo, tam=tam, color=GRIS, interlineado=1.28,
                   color_fuerte=acento)
            cursor = cursor + Inches(alto + 0.16)

    if clave:
        # La caja va justo debajo del contenido, no anclada al pie: así el aire que
        # sobra queda al final de la slide en vez de abrir un hueco en el medio.
        alto_contenido = max(_alto_columna(tam, g) for g in normalizados)
        tope = min(y + Inches(alto_contenido + 0.18),
                   ALTO - Inches(0.62 + alto_clave))
        _caja_titulada(slide, MARGEN, tope, UTIL, Inches(alto_clave - 0.1),
                       clave[0], clave[1], color=acento, tam=12.5,
                       relleno=SUPERFICIE)
    _pie(slide, numero, datos.get("seccion", ""))
    return slide


def slide_ficha(prs, datos, numero):
    """Ficha de indicador: metadata + qué mide + cómo está construido."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    acento = datos.get("color", AZUL)
    y = _encabezado(slide, datos["titulo"], datos.get("bajada"), acento,
                    datos.get("etiqueta", "ficha del indicador"))

    metas = datos["meta"]
    ancho_chip = (UTIL - Inches(0.3) * (len(metas) - 1)) / len(metas)
    for indice, (etiqueta, valor) in enumerate(metas):
        _chip(slide, MARGEN + indice * (ancho_chip + Inches(0.3)), y, etiqueta,
              valor, ancho_chip, acento)
    y = y + Inches(0.88)

    bloques = datos["bloques"]
    disponible = _pulg(ALTO - y) - 0.62
    ancho_bloque = (UTIL - Inches(0.3) * (len(bloques) - 1)) / len(bloques)
    ancho_in = _pulg(ancho_bloque)

    tam = _tam_que_cabe(
        lambda t: max(_alto_caja(lineas, ancho_in, t) for _, lineas in bloques),
        disponible, 12.0, 9.0)
    alto = min(disponible,
               max(_alto_caja(lineas, ancho_in, tam) for _, lineas in bloques))

    for indice, (titulo, lineas) in enumerate(bloques):
        _caja_titulada(slide, MARGEN + indice * (ancho_bloque + Inches(0.3)), y,
                       ancho_bloque, Inches(alto), titulo, lineas,
                       color=acento if indice == 0 else GRIS, tam=tam)
    _pie(slide, numero, datos.get("seccion", ""))
    return slide


def slide_cajas(prs, datos, numero):
    """Tres (o N) cajas paralelas — escenarios, casos, contrastes."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    acento = datos.get("color", VERDE)
    y = _encabezado(slide, datos["titulo"], datos.get("bajada"), acento,
                    datos.get("etiqueta"))

    nota = datos.get("nota")
    alto_nota = (_alto_caja(nota[1], _pulg(UTIL), 12) + 0.1) if nota else 0.0
    cajas = datos["cajas"]
    disponible = _pulg(ALTO - y) - 0.62 - alto_nota
    ancho = (UTIL - Inches(0.28) * (len(cajas) - 1)) / len(cajas)
    ancho_texto = _pulg(ancho) - 0.4

    # El título va a 3,5 pt sobre el cuerpo; medirlo aparte es lo que evita que el
    # párrafo siguiente se le monte cuando ocupa dos líneas. Se reserva el alto del
    # título MÁS largo para todas las cajas, así los cuerpos arrancan alineados.
    def _alto_titulo(tam: float) -> float:
        return max(_alto_texto(c["titulo"], ancho_texto, tam + 3.5, 1.08,
                               negrita=True) for c in cajas)

    def _alto_contenido(tam: float, caja) -> float:
        return (0.52 + _alto_titulo(tam) + 0.1
                + _alto_texto(caja["lineas"], ancho_texto, tam, 1.3) + 0.22)

    tam = _tam_que_cabe(
        lambda t: max(_alto_contenido(t, c) for c in cajas), disponible, 11.5, 8.5)
    alto = min(disponible, max(_alto_contenido(tam, c) for c in cajas))
    alto_titulo = _alto_titulo(tam)

    for indice, caja in enumerate(cajas):
        left = MARGEN + indice * (ancho + Inches(0.28))
        color = caja.get("color", GRIS)
        _rect(slide, left, y, ancho, Inches(alto), relleno=SUPERFICIE)
        _rect(slide, left, y, ancho, Emu(38100), relleno=color)
        cursor = y + Inches(0.22)
        _texto(slide, left + Inches(0.2), cursor, ancho - Inches(0.4), Inches(0.24),
               caja["kicker"].upper(), tam=9.5, color=color, negrita=True)
        cursor = cursor + Inches(0.3)
        _texto(slide, left + Inches(0.2), cursor, ancho - Inches(0.4),
               Inches(alto_titulo), caja["titulo"], tam=tam + 3.5, color=BLANCO,
               negrita=True, interlineado=1.08)
        cursor = cursor + Inches(alto_titulo + 0.1)
        _texto(slide, left + Inches(0.2), cursor, ancho - Inches(0.4),
               Inches(max(0.3, alto - _pulg(cursor - y) - 0.18)), caja["lineas"],
               tam=tam, color=GRIS, interlineado=1.3, espacio_antes=6,
               color_fuerte=color)

    if nota:
        tope = min(y + Inches(alto + 0.22), ALTO - Inches(0.62 + alto_nota))
        _caja_titulada(slide, MARGEN, tope, UTIL, Inches(alto_nota - 0.1),
                       nota[0], nota[1], color=AMBAR, tam=12)
    _pie(slide, numero, datos.get("seccion", ""))
    return slide


def slide_tabla(prs, datos, numero):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    acento = datos.get("color", VERDE)
    y = _encabezado(slide, datos["titulo"], datos.get("bajada"), acento,
                    datos.get("etiqueta"))

    nota = datos.get("nota")
    alto_nota = (_alto_caja(nota[1], _pulg(UTIL), 12) + 0.1) if nota else 0.0
    filas = datos["filas"]
    encabezados = datos["encabezados"]
    disponible = _pulg(ALTO - y) - 0.62 - alto_nota

    pesos = datos.get("pesos", [1] * len(encabezados))
    suma = sum(pesos)
    anchos = [_pulg(UTIL) * peso / suma for peso in pesos]

    # PowerPoint ignora `row.height` si el texto no cabe: expande la fila y la tabla
    # se sale de la slide. Por eso el tamaño se busca hasta que el total quepa.
    INTERLINEADO = 1.16
    MARGEN_V = 0.06

    def _alto_fila(tam: float, fila, es_encabezado: bool) -> float:
        lineas = max(
            _n_lineas(celda, anchos[col] - 0.26, tam,
                      negrita=es_encabezado or col == 0)
            for col, celda in enumerate(fila))
        return lineas * tam * INTERLINEADO / 72.0 + 2 * MARGEN_V

    def _alto_total(tam: float) -> float:
        return (_alto_fila(9.5, encabezados, True)
                + sum(_alto_fila(tam, fila, False) for fila in filas))

    tam = _tam_que_cabe(_alto_total, disponible, 11.0, 8.0)

    forma = slide.shapes.add_table(len(filas) + 1, len(encabezados), MARGEN, y,
                                   UTIL, Inches(_alto_total(tam)))
    tabla = forma.table
    tabla.first_row = False
    tabla.horz_banding = False

    for indice, ancho in enumerate(anchos):
        tabla.columns[indice].width = Inches(ancho)

    def _celda(celda, texto, tam_celda, color, negrita, relleno, tipo=TIPO):
        celda.fill.solid()
        celda.fill.fore_color.rgb = relleno
        celda.margin_left = celda.margin_right = Inches(0.13)
        celda.margin_top = celda.margin_bottom = Inches(MARGEN_V)
        celda.vertical_anchor = MSO_ANCHOR.MIDDLE
        marco = celda.text_frame
        marco.word_wrap = True
        parrafo = marco.paragraphs[0]
        parrafo.line_spacing = INTERLINEADO
        _runs(parrafo, texto, tam_celda, color, tipo, negrita_base=negrita)

    tabla.rows[0].height = Inches(_alto_fila(9.5, encabezados, True))
    for columna, titulo in enumerate(encabezados):
        _celda(tabla.cell(0, columna), titulo.upper(), 9.5, acento, True,
               SUPERFICIE_ALT)

    for indice_fila, fila in enumerate(filas, start=1):
        tabla.rows[indice_fila].height = Inches(_alto_fila(tam, fila, False))
        relleno = SUPERFICIE if indice_fila % 2 else FONDO
        for columna, valor in enumerate(fila):
            color = BLANCO if columna == 0 else GRIS
            _celda(tabla.cell(indice_fila, columna), valor, tam,
                   color, columna == 0, relleno)

    if nota:
        tope = min(y + Inches(_alto_total(tam) + 0.22),
                   ALTO - Inches(0.62 + alto_nota))
        _caja_titulada(slide, MARGEN, tope, UTIL, Inches(alto_nota - 0.1),
                       nota[0], nota[1], color=AMBAR, tam=12)
    _pie(slide, numero, datos.get("seccion", ""))
    return slide


def slide_ejercicio(prs, datos, numero):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    y = _encabezado(slide, datos["titulo"], None, AMBAR,
                    datos.get("etiqueta", "ejercicio de interpretación"))

    ancho_izq = Inches(4.3)
    disponible = _pulg(ALTO - y) - 0.62
    left = MARGEN + ancho_izq + Inches(0.3)
    ancho_der = UTIL - ancho_izq - Inches(0.3)
    izq_in, der_in = _pulg(ancho_izq), _pulg(ancho_der)
    HUECO = 0.2

    tam_dato = _tam_que_cabe(
        lambda t: _alto_caja(datos["dato"], izq_in, t, 1.24), disponible, 12.0, 8.5)
    _caja_titulada(slide, MARGEN, y, ancho_izq, Inches(disponible),
                   "el dato publicado", datos["dato"], color=AMBAR, tam=tam_dato)

    # Las dos cajas de la derecha comparten la altura: se busca un tamaño con el que
    # ambas quepan y luego cada una toma exactamente lo que su contenido necesita.
    tam = _tam_que_cabe(
        lambda t: (_alto_caja(datos["lectura"], der_in, t)
                   + _alto_caja(datos["precio"], der_in, t) + HUECO),
        disponible, 12.0, 8.5)

    alto_lectura = _alto_caja(datos["lectura"], der_in, tam)
    alto_precio = max(_alto_caja(datos["precio"], der_in, tam),
                      disponible - alto_lectura - HUECO)
    _caja_titulada(slide, left, y, ancho_der, Inches(alto_lectura),
                   "cómo se interpreta", datos["lectura"], color=VERDE, tam=tam)
    _caja_titulada(slide, left, y + Inches(alto_lectura + HUECO), ancho_der,
                   Inches(alto_precio),
                   "qué debería hacer el precio, y por qué", datos["precio"],
                   color=AZUL, tam=tam)
    _pie(slide, numero, datos.get("seccion", ""))
    return slide


def slide_cierre(prs, datos, numero):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo(slide)
    _rect(slide, Emu(0), Emu(0), Inches(0.14), ALTO, relleno=VERDE)
    _texto(slide, MARGEN, Inches(0.95), UTIL, Inches(0.3), datos["kicker"].upper(),
           tam=11.5, color=VERDE, negrita=True)
    tope_titulo, TAM_TITULO = 1.32, 36
    alto_titulo = _alto_texto(datos["titulo"], 10.0, TAM_TITULO, 1.0, negrita=True)
    _texto(slide, MARGEN, Inches(tope_titulo), Inches(10.0), Inches(alto_titulo),
           datos["titulo"], tam=TAM_TITULO, color=BLANCO, negrita=True,
           interlineado=1.0)

    ancho_in = _pulg(UTIL)
    alto_aviso = _alto_caja(datos["aviso"][1], ancho_in, 12)
    alto_proceso = _alto_caja(datos["proceso"][1], ancho_in, 12)

    # El bloque arranca donde termina el título medido, no en una constante: con una
    # constante el filete de la advertencia acababa cruzando el texto del título.
    y = Inches(tope_titulo + alto_titulo + 0.26)
    _caja_titulada(slide, MARGEN, y, UTIL, Inches(alto_aviso),
                   datos["aviso"][0], datos["aviso"][1], color=AMBAR, tam=12)
    y = y + Inches(alto_aviso + 0.16)
    _caja_titulada(slide, MARGEN, y, UTIL, Inches(alto_proceso),
                   datos["proceso"][0], datos["proceso"][1], color=AZUL, tam=12)
    y = y + Inches(alto_proceso + 0.2)
    _bloque_autor(slide, MARGEN, y, datos["autor"], FOTO, diametro=1.02,
                  tam_nombre=13, compacto=True)
    _texto(slide, MARGEN + Inches(9.35), y + Inches(0.3), Inches(2.6), Inches(0.24),
           datos["fecha"], tam=11, color=GRIS)
    return slide


RENDER = {
    "portada": slide_portada,
    "seccion": slide_seccion,
    "bullets": slide_bullets,
    "ficha": slide_ficha,
    "cajas": slide_cajas,
    "tabla": slide_tabla,
    "ejercicio": slide_ejercicio,
    "cierre": slide_cierre,
}


def construir(slides, destino: Path, foto: Path | None = None,
              zoom: float = 1.0, centro_y: float = 0.5) -> Path:
    global FOTO
    origen = _resolver_foto(foto)
    if origen is None:
        print("  aviso: sin retrato disponible, la maqueta va sin foto")
    FOTO = _foto_circular(origen, zoom=zoom, centro_y=centro_y) if origen else None

    prs = Presentation()
    prs.slide_width = ANCHO
    prs.slide_height = ALTO
    for indice, datos in enumerate(slides, start=1):
        RENDER[datos["tipo"]](prs, datos, indice)
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destino))
    return destino


def main() -> None:
    from capacitacion_fundamental_contenido import SLIDES  # noqa: PLC0415

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        default="docs/capacitacion/Analisis Fundamental - Capacitacion GI.pptx",
        help="ruta del .pptx de salida",
    )
    parser.add_argument(
        "--foto",
        default="docs/capacitacion/assets/autor.png",
        help="retrato del autor; si la ruta no existe se busca otra imagen en la "
             "misma carpeta. Se recorta en círculo",
    )
    # 1.15 / 0.47 se eligieron comparando recortes: da presencia al rostro y aún deja
    # ver el traje y el nudo de la corbata. Con 1.3 o más, el traje se pierde y la
    # cabeza queda flotando dentro del círculo.
    parser.add_argument(
        "--zoom", type=float, default=1.15,
        help="cierra el encuadre sobre el rostro (1.0 = círculo inscrito completo)",
    )
    parser.add_argument(
        "--centro", type=float, default=0.47,
        help="fracción de la altura donde queda el centro del recorte",
    )
    args = parser.parse_args()
    foto = Path(args.foto) if args.foto else None
    ruta = construir(SLIDES, Path(args.out), foto, args.zoom, args.centro)
    print(f"{len(SLIDES)} slides -> {ruta}")


if __name__ == "__main__":
    import sys

    sys.path.insert(0, str(Path(__file__).parent))
    main()
