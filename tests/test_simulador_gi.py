# -*- coding: utf-8 -*-
"""Pruebas de integridad y especificaciones para Simulador GI Real.xlsx."""
from pathlib import Path
import openpyxl

REPO = Path(__file__).resolve().parent.parent
EXCEL_DIR = REPO / "calculadoras excel"
SIMULADOR_REAL = EXCEL_DIR / "Simulador GI Real.xlsx"
SIMULADOR_ALIAS = EXCEL_DIR / "Simulador GI.xlsx"

GRUPOS_CANONICOS = {
    "FX",
    "CFD Commodities",
    "CFD Indices",
    "Criptomonedas",
    "CFD ETF",
    "CFD Acciones",
}

COLUMNAS_DATOS_ESPERADAS = [
    "INSTRUMENTO",
    "TICKER MT5",
    "DECIMALES",
    "TAMAÑO DE CONTRATO (1 lote)",
    "VALOR EN CLP DE 1,0 DE PRECIO (POR LOTE)",
    "MARGEN",
    "SPREAD",
    "PRECIO DE REFERENCIA",
    "MODO SWAP",
    "SWAP COMPRA",
    "SWAP VENTA",
    "VOLUMEN MÍNIMO",
    "PASO DE VOLUMEN",
    "GRUPO",
]


def test_archivos_simulador_existen():
    """Verifica que ambos archivos Excel existen y tienen tamaño coherente."""
    assert SIMULADOR_REAL.is_file(), f"Falta {SIMULADOR_REAL}"
    assert SIMULADOR_REAL.stat().st_size > 15_000, "Simulador GI Real.xlsx parece truncado o vacío"
    assert SIMULADOR_ALIAS.is_file(), f"Falta {SIMULADOR_ALIAS}"
    assert SIMULADOR_ALIAS.stat().st_size > 15_000, "Simulador GI.xlsx parece truncado o vacío"


def test_hojas_y_estructura_general():
    """Valida la carga del libro y que existan las hojas Simulador y Datos."""
    wb = openpyxl.load_workbook(SIMULADOR_REAL, data_only=False)
    assert "Simulador" in wb.sheetnames, "Falta la hoja 'Simulador'"
    assert "Datos" in wb.sheetnames, "Falta la hoja 'Datos'"

    ws_sim = wb["Simulador"]
    assert ws_sim["C1"].value == "SIMULADOR DE OPERACIONES"
    assert ws_sim["F4"].value is not None, "El instrumento seleccionado por defecto en F4 no debe ser None"
    assert ws_sim["F5"].value == 2000000, "El capital inicial de ejemplo en F5 debe ser $2.000.000"
    assert ws_sim.protection.sheet is True, "La hoja Simulador debe estar protegida para evitar sobreescritura accidental"


def test_hoja_datos_columnas_e_integridad():
    """Verifica exhaustivamente el catálogo de datos y especificaciones de contrato."""
    wb = openpyxl.load_workbook(SIMULADOR_REAL, data_only=False)
    ws_datos = wb["Datos"]

    headers = [ws_datos.cell(1, col).value for col in range(1, len(COLUMNAS_DATOS_ESPERADAS) + 1)]
    assert headers == COLUMNAS_DATOS_ESPERADAS, f"Encabezados no coinciden: {headers}"

    filas_datos = 0
    tickers_vistos = set()
    for row in range(2, ws_datos.max_row):
        instrumento = ws_datos.cell(row, 1).value
        ticker = ws_datos.cell(row, 2).value
        if not ticker or not instrumento:
            continue

        filas_datos += 1
        tickers_vistos.add(ticker)

        digits = ws_datos.cell(row, 3).value
        contrato = ws_datos.cell(row, 4).value
        clp_unidad = ws_datos.cell(row, 5).value
        margen = ws_datos.cell(row, 6).value
        spread = ws_datos.cell(row, 7).value
        precio_ref = ws_datos.cell(row, 8).value
        swap_modo = ws_datos.cell(row, 9).value
        swap_compra = ws_datos.cell(row, 10).value
        swap_venta = ws_datos.cell(row, 11).value
        vol_min = ws_datos.cell(row, 12).value
        vol_paso = ws_datos.cell(row, 13).value
        grupo = ws_datos.cell(row, 14).value

        assert isinstance(digits, int) and 0 <= digits <= 6, f"Digits inválido en {ticker}: {digits}"
        assert isinstance(contrato, (int, float)) and contrato > 0, f"Contrato inválido en {ticker}: {contrato}"
        assert isinstance(clp_unidad, (int, float)) and clp_unidad > 0, f"CLP unidad inválido en {ticker}: {clp_unidad}"
        assert isinstance(margen, (int, float)) and 0.0001 <= margen <= 1.0, f"Margen inválido en {ticker}: {margen}"
        assert isinstance(spread, (int, float)) and spread >= 0, f"Spread inválido en {ticker}: {spread}"
        assert isinstance(precio_ref, (int, float)) and precio_ref > 0, f"Precio de ref inválido en {ticker}: {precio_ref}"
        assert swap_modo in (1, 5), f"Modo swap inválido en {ticker}: {swap_modo} (debe ser 1 o 5)"
        assert isinstance(swap_compra, (int, float)), f"Swap compra inválido en {ticker}: {swap_compra}"
        assert isinstance(swap_venta, (int, float)), f"Swap venta inválido en {ticker}: {swap_venta}"
        assert isinstance(vol_min, (int, float)) and vol_min > 0, f"Vol min inválido en {ticker}: {vol_min}"
        assert isinstance(vol_paso, (int, float)) and vol_paso > 0, f"Vol paso inválido en {ticker}: {vol_paso}"
        assert grupo in GRUPOS_CANONICOS, f"Grupo desconocido en {ticker}: {grupo}"

    assert filas_datos >= 140, f"Se esperaban al menos 140 instrumentos operables, encontrados: {filas_datos}"


def test_activos_clave_especificaciones():
    """Valida especificaciones de activos estratégicos según reglas de negocio."""
    wb = openpyxl.load_workbook(SIMULADOR_REAL, data_only=False)
    ws_datos = wb["Datos"]

    activos = {}
    for r in range(2, ws_datos.max_row):
        t = ws_datos.cell(r, 2).value
        if t:
            activos[t] = {
                "nombre": ws_datos.cell(r, 1).value,
                "contrato": ws_datos.cell(r, 4).value,
                "digits": ws_datos.cell(r, 3).value,
                "swap_modo": ws_datos.cell(r, 9).value,
                "vol_min": ws_datos.cell(r, 12).value,
                "vol_paso": ws_datos.cell(r, 13).value,
                "grupo": ws_datos.cell(r, 14).value,
            }

    # USDCLP
    assert "USDCLP" in activos
    assert activos["USDCLP"]["contrato"] == 100_000
    assert activos["USDCLP"]["digits"] == 2
    assert activos["USDCLP"]["swap_modo"] == 1
    assert activos["USDCLP"]["vol_min"] == 0.01

    # XAUUSD
    assert "XAUUSD" in activos
    assert activos["XAUUSD"]["contrato"] == 100
    assert activos["XAUUSD"]["digits"] == 2
    assert activos["XAUUSD"]["swap_modo"] == 1

    # #AAPL (Acción: 1 título por lote en cuenta real)
    assert "#AAPL" in activos
    assert activos["#AAPL"]["contrato"] == 1
    assert activos["#AAPL"]["swap_modo"] == 5
    assert activos["#AAPL"]["vol_paso"] == 0.1

    # QQQ.US (ETF: 1 título por lote, paso 0.1)
    assert "QQQ.US" in activos
    assert activos["QQQ.US"]["contrato"] == 1
    assert activos["QQQ.US"]["swap_modo"] == 5
    assert activos["QQQ.US"]["vol_paso"] == 0.1

    # BTCUSD (Cripto: contrato 1, modo interés)
    assert "BTCUSD" in activos
    assert activos["BTCUSD"]["contrato"] == 1
    assert activos["BTCUSD"]["swap_modo"] == 5


def test_certificacion_cuenta_real_y_umbrales():
    """Valida que los metadatos certifiquen la cuenta REAL 51492 y los umbrales de riesgo."""
    wb = openpyxl.load_workbook(SIMULADOR_REAL, data_only=False)
    ws_sim = wb["Simulador"]
    ws_datos = wb["Datos"]

    pie = str(ws_sim["C27"].value)
    assert "cuenta real" in pie.lower(), f"El pie debe certificar cuenta real: {pie}"
    assert "200%" in pie, f"La llamada a margen debe ser 200%: {pie}"
    assert "50%" in pie, f"El cierre forzado debe ser 50%: {pie}"

    nota_datos = str(ws_datos.cell(ws_datos.max_row, 1).value)
    assert "51492" in nota_datos, f"La nota de Datos debe citar la cuenta 51492: {nota_datos}"
    assert "REAL" in nota_datos, f"La nota de Datos debe certificar tipo REAL: {nota_datos}"
    assert "CLP" in nota_datos, f"La nota de Datos debe certificar moneda CLP: {nota_datos}"


def test_formulas_clave_simulador():
    """Verifica que las fórmulas críticas de simulación, swap y margen permanezcan íntegras."""
    wb = openpyxl.load_workbook(SIMULADOR_REAL, data_only=False)
    ws_sim = wb["Simulador"]

    # Fila 11 (primera operación)
    f_margen = str(ws_sim["E11"].value)
    assert "$D11" in f_margen and "$G11" in f_margen and "$R$4" in f_margen and "$R$5" in f_margen

    f_apal = str(ws_sim["F11"].value)
    assert "$D11" in f_apal and "$G11" in f_apal and "$R$4/$F$5" in f_apal

    f_bruto = str(ws_sim["J11"].value)
    assert "$H11-$G11" in f_bruto and "$G11-$H11" in f_bruto

    f_spread = str(ws_sim["K11"].value)
    assert "$D11*$R$6*$R$4" in f_spread

    f_swap = str(ws_sim["L11"].value)
    assert "$R$10" in f_swap and "$R$9" in f_swap and "POWER(10,-$R$2)" in f_swap and "360" in f_swap

    f_neto = str(ws_sim["M11"].value)
    assert "N($J11)-N($K11)-N($L11)" in f_neto

    # Cierre
    f_monto_ctrl = str(ws_sim["M19"].value)
    assert "SUMPRODUCT($D$11:$D$16,$G$11:$G$16)*$R$4" in f_monto_ctrl

    f_recorrido = str(ws_sim["M23"].value)
    assert "$R$12/100*$M$21" in f_recorrido, "El recorrido adverso debe medirse contra el cierre forzado $R$12"
